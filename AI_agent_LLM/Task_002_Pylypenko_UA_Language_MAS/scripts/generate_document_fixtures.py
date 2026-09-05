"""Створює невеликі локальні TXT, DOCX, PDF і PNG для відтворюваної демонстрації intake."""

from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFont, PngImagePlugin
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "fixtures" / "documents"
DOCUMENTS = {
    "txt": [
        "НАВЧАЛЬНИЙ НАКАЗ (усі ПІБ вигадані)",
        "На мою думку це даний документ потребує уточнення.",
        "Згідно наказу директора Олексія Коваленко відділ має приймати участь у нараді.",
        "Будьласка, передайте проєкт Ігорю Журавель до встановленого дедлайну.",
        "На даний момент основна суть пропозиції співпадає з попереднім планом.",
        "Цей план план потрібно перевірити та прийняти міри.",
        "Повторюваний фрагмент доручення.",
        "Повторюваний фрагмент доручення.",
        "Форму «для Марії Ковальчук» наведено як правильний контрольний приклад.",
    ],
    "docx": [
        "СЛУЖБОВА ЗАПИСКА (усі ПІБ вигадані)",
        "По питанню закупівлі даний документ являється основним.",
        "Звіт підготовлено Тарасом Бондар та передано Соломії Круть.",
        "Слідуючий етап — заключити договір до п'ятниці.",
        "Будьласка, погодьте дедлайн і прийміть участь в обговоренні.",
        "Спільна співпраця команди має місце на протязі двох місяців.",
        "Рішення рішення потрібно повторно перевірити.",
        "Однаковий підсумковий абзац.",
        "Однаковий підсумковий абзац.",
    ],
    "pdf_page_1": [
        "НАВЧАЛЬНИЙ ПРОТОКОЛ (усі ПІБ вигадані)",
        "На мою думку це вірне рішення комісії.",
        "Згідно наказу слово надали Ігорю Журавель.",
        "По питанню строків доповів Андрій Швець.",
        "Комісія вирішила прийняти міри та заключити договір.",
    ],
    "pdf_page_2": [
        "ПРОДОВЖЕННЯ ПРОТОКОЛУ — сторінка 2",
        "На даний момент пам'ятний сувенір уже підготовлено.",
        "Будьласка, повідомте дедлайн Олексію Коваленко.",
        "Протокол протокол передати секретареві.",
        "Цей висновок повторено дослівно.",
        "Цей висновок повторено дослівно.",
    ],
    "png": [
        "НАВЧАЛЬНЕ ОГОЛОШЕННЯ (усі ПІБ вигадані)",
        "На мою думку це даний документ.",
        "Запрошуємо Андрія Швець на слідуючий етап відбору.",
        "На даний момент відкрита вільна вакансія редактора.",
        "Будьласка, надайте відповідь до дедлайну.",
        "Вірне рішення рішення буде опубліковано завтра.",
        "Форму «від Олени Журавель» залишено як правильний контрольний приклад.",
    ],
}


def font_path() -> Path:
    candidates = [Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/calibri.ttf")]
    return next(path for path in candidates if path.is_file())


def main() -> None:
    TARGET.mkdir(parents=True, exist_ok=True)
    (TARGET / "example.txt").write_text("\n".join(DOCUMENTS["txt"]), encoding="utf-8")

    docx = Document()
    docx.add_heading("Приклад DOCX", level=1)
    for line in DOCUMENTS["docx"]:
        docx.add_paragraph(line)
    docx.save(TARGET / "example.docx")

    pdfmetrics.registerFont(TTFont("FixtureFont", str(font_path())))
    pdf = canvas.Canvas(str(TARGET / "example.pdf"))
    pdf.setFont("FixtureFont", 13)
    for index, line in enumerate(DOCUMENTS["pdf_page_1"]):
        pdf.drawString(60, 800 - index * 30, line)
    pdf.showPage()
    pdf.setFont("FixtureFont", 13)
    for index, line in enumerate(DOCUMENTS["pdf_page_2"]):
        pdf.drawString(60, 800 - index * 30, line)
    pdf.save()

    image_lines = DOCUMENTS["png"]
    image = Image.new("RGB", (1500, 80 + len(image_lines) * 52), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(str(font_path()), 30)
    for index, line in enumerate(image_lines):
        draw.text((30, 30 + index * 52), line, fill="black", font=font)
    metadata = PngImagePlugin.PngInfo()
    metadata.add_text("language_mas_fixture_text", "\n".join(image_lines))
    image.save(TARGET / "example.png", pnginfo=metadata)
    print(f"Створено fixtures: {TARGET}")


if __name__ == "__main__":
    main()
